"""
tests/test_resume.py

Integration tests for the resume upload, retrieval, and deletion endpoints.
"""

import io
import pytest
import fitz
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.models.user import users_collection
from app.models.resume import resumes_collection

client = TestClient(app)

TEST_EMAIL = "testuser_resume@example.com"
TEST_PASSWORD = "supersecurepassword123"
TEST_EMAIL_2 = "testuser_resume2@example.com"

def get_auth_token(email=TEST_EMAIL):
    """Helper to get a valid token."""
    response = client.post(
        "/api/auth/login",
        data={"username": email, "password": TEST_PASSWORD}
    )
    return response.json()["access_token"]


@pytest.fixture(autouse=True)
def setup_teardown():
    """Create test users and clean up database before and after each test."""
    users_collection.delete_many({"email": {"$in": [TEST_EMAIL, TEST_EMAIL_2]}})
    resumes_collection.delete_many({})
    
    # Register user 1
    client.post(
        "/api/auth/register",
        json={"name": "Test User 1", "email": TEST_EMAIL, "password": TEST_PASSWORD}
    )
    # Register user 2
    client.post(
        "/api/auth/register",
        json={"name": "Test User 2", "email": TEST_EMAIL_2, "password": TEST_PASSWORD}
    )
    
    yield
    
    users_collection.delete_many({"email": {"$in": [TEST_EMAIL, TEST_EMAIL_2]}})
    resumes_collection.delete_many({})


def create_dummy_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(fitz.Point(50, 50), "This is a test PDF resume.\nHas multiple lines.")
    return doc.write()


def create_dummy_docx() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("This is a test DOCX resume.")
    doc.add_paragraph("Has multiple lines.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_upload_pdf():
    token = get_auth_token()
    pdf_bytes = create_dummy_pdf()
    
    response = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test_resume.pdf", pdf_bytes, "application/pdf")}
    )
    
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "test_resume.pdf"
    assert data["file_type"] == "pdf"
    assert data["file_size"] > 0
    assert "id" in data  # Fixed: API now returns "id" not "_id"


def test_upload_docx():
    token = get_auth_token()
    docx_bytes = create_dummy_docx()
    
    response = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    
    assert response.status_code == 201
    data = response.json()
    assert data["filename"] == "test_resume.docx"
    assert data["file_type"] == "docx"
    assert data["file_size"] > 0
    assert "id" in data  # Fixed: API now returns "id" not "_id"


def test_unauthenticated_upload():
    pdf_bytes = create_dummy_pdf()
    response = client.post(
        "/api/resumes/upload",
        files={"file": ("test.pdf", pdf_bytes, "application/pdf")}
    )
    assert response.status_code == 401


def test_unsupported_file_type():
    token = get_auth_token()
    response = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test_resume.txt", b"Hello text file", "text/plain")}
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_oversized_file():
    token = get_auth_token()
    # Create a 6MB dummy file
    large_bytes = b"0" * (6 * 1024 * 1024)
    response = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("large.pdf", large_bytes, "application/pdf")}
    )
    assert response.status_code == 413
    assert "too large" in response.json()["detail"].lower()


def test_empty_file():
    token = get_auth_token()
    response = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("empty.pdf", b"", "application/pdf")}
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_list_resumes():
    token = get_auth_token()
    
    # Upload 2 resumes
    pdf_bytes = create_dummy_pdf()
    client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test1.pdf", pdf_bytes, "application/pdf")}
    )
    client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test2.pdf", pdf_bytes, "application/pdf")}
    )
    
    response = client.get(
        "/api/resumes",
        headers={"Authorization": f"Bearer {token}"}
    )
    
    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 2
    assert len(data["resumes"]) == 2
    # Ensure newest first (test2 was uploaded second)
    assert data["resumes"][0]["filename"] == "test2.pdf"
    assert "extracted_text" not in data["resumes"][0]


def test_get_resume_with_text():
    token = get_auth_token()
    pdf_bytes = create_dummy_pdf()
    
    upload_res = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test1.pdf", pdf_bytes, "application/pdf")}
    )
    resume_id = upload_res.json()["id"]  # Fixed: "id" not "_id"
    
    response = client.get(
        f"/api/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {token}"}
    )
    
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "test1.pdf"
    assert "This is a test PDF resume." in data["extracted_text"]


def test_unauthorized_resume_access():
    # User 1 uploads
    token1 = get_auth_token(TEST_EMAIL)
    upload_res = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token1}"},
        files={"file": ("test1.pdf", create_dummy_pdf(), "application/pdf")}
    )
    resume_id = upload_res.json()["id"]  # Fixed: "id" not "_id"
    
    # User 2 tries to access
    token2 = get_auth_token(TEST_EMAIL_2)
    response = client.get(
        f"/api/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {token2}"}
    )
    assert response.status_code == 404


def test_delete_resume():
    token = get_auth_token()
    upload_res = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("test1.pdf", create_dummy_pdf(), "application/pdf")}
    )
    resume_id = upload_res.json()["id"]  # Fixed: "id" not "_id"
    
    # Delete
    del_res = client.delete(
        f"/api/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert del_res.status_code == 204
    
    # Verify it's gone
    get_res = client.get(
        f"/api/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert get_res.status_code == 404


def test_unauthorized_resume_deletion():
    # User 1 uploads
    token1 = get_auth_token(TEST_EMAIL)
    upload_res = client.post(
        "/api/resumes/upload",
        headers={"Authorization": f"Bearer {token1}"},
        files={"file": ("test1.pdf", create_dummy_pdf(), "application/pdf")}
    )
    resume_id = upload_res.json()["id"]  # Fixed: "id" not "_id"
    
    # User 2 tries to delete
    token2 = get_auth_token(TEST_EMAIL_2)
    del_res = client.delete(
        f"/api/resumes/{resume_id}",
        headers={"Authorization": f"Bearer {token2}"}
    )
    assert del_res.status_code == 404
